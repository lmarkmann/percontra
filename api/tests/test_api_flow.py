import json
from io import BytesIO

import pytest
from django.test import Client
from openpyxl import load_workbook

from percontra import submissions


@pytest.fixture
def operator(tmp_path, settings):
    settings.MIGRATION_DB = str(tmp_path / "migration.duckdb")
    client = Client(enforce_csrf_checks=True)
    token = client.get("/api/csrf").json()["token"]

    def post(path, payload):
        return client.post(
            path, json.dumps(payload), content_type="application/json", HTTP_X_CSRFTOKEN=token
        )

    return client, post


def test_http_review_release_export_evidence_and_invalidation(operator):
    client, post = operator
    loaded = post("/api/example", {})
    assert loaded.status_code == 200
    batches = loaded.json()["batches"]
    gap_batch = next(row["id"] for row in batches if row["key"]["source_batch_id"] == "public-gap")
    clean_batch = next(
        row["id"] for row in batches if row["key"]["source_batch_id"] == "public-clean"
    )
    assert (
        post(f"/api/releases/{gap_batch}/approve", {"author": "API test reviewer"}).status_code
        == 400
    )
    assert (
        post(f"/api/releases/{clean_batch}/approve", {"author": "API test reviewer"}).status_code
        == 200
    )
    decision = {
        "gap_row": 2,
        "target_row": 3,
        "author": "API test reviewer",
        "reason": "Synthetic expense treatment",
    }
    assert post("/api/decisions", decision).status_code == 200
    postings = client.get("/api/postings", {"batch": gap_batch}).json()
    assert postings["total"] == 2
    resolved = next(row for row in postings["items"] if row["decisions_used"])
    assert resolved["destination"]["investor_amount_local"]["amount"] == "76.25"
    approved = post(f"/api/releases/{gap_batch}/approve", {"author": "API test reviewer"})
    assert approved.json()["state"] == "approved"
    exported = post("/api/exports", {"batch": gap_batch})
    assert exported.status_code == 200
    assert exported["X-Export-Status"] == "Exported; destination not checked"
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert workbook.active.max_column == 27
    assert workbook.active.max_row == 3
    workbook.close()
    evidence = client.get(f"/api/postings/{resolved['posting_id']}/evidence").json()
    assert evidence["source"][0]["source"]["physical_row"] == 4
    assert evidence["decisions"][0]["author"] == "API test reviewer"
    assert evidence["approval"]["state"] == "approved"
    assert (
        post("/api/decisions", {**decision, "reason": "New review, same target"}).status_code == 200
    )
    assert post("/api/exports", {"batch": gap_batch}).status_code == 400
    overview = client.get("/api/overview").json()
    states = {row["id"]: row["release"]["state"] for row in overview["batches"]}
    assert states == {gap_batch: "stale", clean_batch: "approved"}
    history = client.get(f"/api/postings/{resolved['posting_id']}/evidence").json()
    assert [row["version"] for row in history["decisions"]] == [1, 2]


def test_submission_requires_explicit_company_confirmation(operator, monkeypatch):
    _, post = operator
    calls = []
    monkeypatch.setattr(submissions, "submit", lambda *args: calls.append(args))
    assert post("/api/submissions", {"batch": "unknown"}).status_code == 400
    assert calls == []


def test_invalid_decision_never_appends_history(operator):
    client, post = operator
    assert post("/api/example", {}).status_code == 200
    rejected = post(
        "/api/decisions",
        {
            "gap_row": 2,
            "target_row": 2,
            "author": "API reviewer",
            "reason": "Asset is not an expense",
        },
    )
    assert rejected.status_code == 400
    assert client.get("/api/overview").json()["gaps"][0]["history"] == []
