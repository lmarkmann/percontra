import ast
from io import BytesIO
from pathlib import Path

import pytest
from conftest import batch
from django.test import Client
from openpyxl import Workbook, load_workbook

from percontra.adapters.base import UploadedFile
from percontra.adapters.destination.erpnext import (
    COMPANY,
    SITE,
    Connection,
    ERPError,
    ERPNextAdapter,
)
from percontra.adapters.registry import catalog, create
from percontra.adapters.source.csv_generic import CSVAdapter
from percontra.ingest.xlsx import read_sheet


def test_an_incomplete_batch_cannot_be_approved(migration):
    gap = batch(migration, "public-gap")
    with pytest.raises(ValueError, match="complete batch"):
        migration.approve(gap, "Demo reviewer")


def test_export_is_labelled_unchecked_and_carries_the_batch_rows(migration):
    gap = batch(migration, "public-gap")
    migration.decide(2, 3, "Demo reviewer", "Administration expense, confirmed in source")
    migration.approve(gap, "Demo reviewer")
    artifact, record = migration.export(gap)
    assert record.label == "Exported; destination not checked"
    workbook = load_workbook(BytesIO(artifact.content), read_only=True)
    assert workbook.active.max_row == 3
    workbook.close()


def test_evidence_keeps_mapping_and_decision_history(migration):
    migration.decide(2, 3, "Demo reviewer", "First decision")
    migration.decide(2, 4, "Demo reviewer", "Second decision")
    posting = next(row for row in migration.postings() if row.decisions_used)
    evidence = migration.evidence(posting.posting_id)
    assert len(evidence["decisions"]) == 2
    assert evidence["source"][0]["source"]["physical_row"] == 4
    assert evidence["mappings"]


def test_row_numbers_survive_blank_rows_and_second_row_header(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LE Mapping"
    sheet.append(["Source", None, "Destination"])
    sheet.append(["Legal Entity", "Currency", "Corvus LE"])
    sheet.append(["Example", "USD", "Example"])
    sheet.append([None, None, None])
    sheet.append(["Other", "USD", "Other"])
    path = tmp_path / "headers.xlsx"
    workbook.save(path)
    parsed = read_sheet(path, "LE Mapping")
    assert parsed.frame["physical_row"].to_list() == [3, 5]
    assert parsed.headers[0] == "Legal Entity"


def test_csv_keeps_start_line_with_blank_and_multiline_records(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text(
        'entity,amount,note\nExample,125.50,"first\nsecond"\n\nExample,-125.50,offset\n'
    )
    adapter = CSVAdapter({"Legal_Entity": "entity", "Amount_(Local_Currency)": "amount"}, {})
    handle = adapter.open(files=[UploadedFile(path, "gl")])
    assert [row.source.physical_row for row in handle.rows] == [2, 5]
    assert handle.rows[0].values["Amount_(Local_Currency)"] == "125.50"


def test_capabilities_and_unverified_adapters_are_explicit():
    for adapter in catalog():
        assert set(adapter["capabilities"].values()) <= {
            "verified",
            "partial",
            "unverified",
            "not_applicable",
        }
    with pytest.raises(NotImplementedError, match="no verified integration"):
        create("entrilia").render(postings=[])


def test_core_has_no_adapter_or_persistence_imports():
    root = Path(__file__).parents[1] / "percontra/core"
    for path in root.glob("*.py"):
        for node in ast.walk(ast.parse(path.read_text())):
            if isinstance(node, ast.ImportFrom):
                assert not any(
                    word in (node.module or "")
                    for word in ("adapters", "store", "django", "duckdb")
                )


def test_erpnext_rejects_other_companies_before_writes(migration):
    adapter = ERPNextAdapter(Connection(SITE, "test-key", "test-secret"))
    problems = adapter.validate(
        postings=migration.batch(batch(migration, "public-clean")), registries={"problems": []}
    )
    assert "company" in {item.code for item in problems}
    with pytest.raises(ERPError, match="disabled"):
        adapter.request("POST", "/api/resource/Journal Entry", body={})


def test_erpnext_checks_ledger_by_account_and_currency():
    adapter = ERPNextAdapter(Connection(SITE, "test-key", "test-secret"))
    document = {
        "name": "TEST-JV",
        "accounts": [
            {
                "account": "Cash",
                "debit_in_account_currency": "125.50",
                "credit_in_account_currency": "0",
            },
            {
                "account": "Income",
                "debit_in_account_currency": "0",
                "credit_in_account_currency": "125.50",
            },
        ],
    }
    adapter.records = lambda *args: [
        {
            "company": COMPANY,
            "account": "Cash",
            "account_currency": "USD",
            "debit": "125.50",
            "credit": "0",
            "is_cancelled": 0,
        },
        {
            "company": COMPANY,
            "account": "Income",
            "account_currency": "USD",
            "debit": "0",
            "credit": "125.50",
            "is_cancelled": 0,
        },
    ]
    assert len(adapter.verify_ledger(document)) == 2
    document["accounts"][0]["debit_in_account_currency"] = "125.51"
    with pytest.raises(ERPError, match="totals"):
        adapter.verify_ledger(document)


def test_csrf_and_nonlocal_writes_are_blocked(tmp_path, settings):
    settings.MIGRATION_DB = str(tmp_path / "api.duckdb")
    client = Client(enforce_csrf_checks=True)
    assert client.post("/api/example").status_code == 403
    token = client.get("/api/csrf").json()["token"]
    assert (
        client.post("/api/example", HTTP_X_CSRFTOKEN=token, REMOTE_ADDR="10.0.0.3").status_code
        == 400
    )
    response = client.post("/api/example", HTTP_X_CSRFTOKEN=token, REMOTE_ADDR="127.0.0.1")
    assert response.status_code == 200
    assert response.json()["loaded"]
